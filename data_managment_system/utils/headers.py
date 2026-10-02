import csv
from enum import StrEnum
from io import BytesIO, StringIO

from openpyxl import load_workbook


class HeaderErrorReason(StrEnum):
    UNREADABLE = "unreadable"
    ENCODING = "encoding"
    BLANK_HEADER = "blank_header"
    NO_COLUMNS = "no_columns"


class HeaderError(Exception):
    def __init__(self, reason: HeaderErrorReason) -> None:
        self.reason = reason
        super().__init__(reason.value)


def _validate_headers(headers: list[str]) -> list[str]:
    while headers and headers[-1] == "":
        headers.pop()
    if not headers:
        raise HeaderError(HeaderErrorReason.NO_COLUMNS)
    if any(header.strip() == "" for header in headers):
        raise HeaderError(HeaderErrorReason.BLANK_HEADER)
    return headers


def _csv_headers(content: bytes) -> list[str]:
    encoding = (
        "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    )
    try:
        text = content.decode(encoding)
    except UnicodeDecodeError as exc:
        raise HeaderError(HeaderErrorReason.ENCODING) from exc
    return _validate_headers(next(csv.reader(StringIO(text)), []))


def _xlsx_headers(content: bytes) -> list[str]:
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            raise HeaderError(HeaderErrorReason.NO_COLUMNS)
        row = next(sheet.iter_rows(values_only=True), ())
        return _validate_headers(
            [cell if isinstance(cell, str) else str(cell or "") for cell in row]
        )
    except HeaderError:
        raise
    except Exception as exc:
        raise HeaderError(HeaderErrorReason.UNREADABLE) from exc
    finally:
        if workbook is not None:
            workbook.close()


def extract_csv_excel_headers(filename: str, content: bytes) -> list[str]:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        return _csv_headers(content)
    if suffix == "xlsx":
        return _xlsx_headers(content)
    raise HeaderError(HeaderErrorReason.UNREADABLE)


__all__ = ["HeaderError", "HeaderErrorReason", "extract_csv_excel_headers"]
