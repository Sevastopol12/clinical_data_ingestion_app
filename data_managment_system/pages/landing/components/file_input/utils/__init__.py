import csv
import logging
from io import StringIO, BytesIO
from openpyxl import load_workbook
from asyncio import Semaphore, gather, create_task
from httpx import AsyncClient

from data_managment_system.models import (
    IngestionResponse,
    IngestionCreate,
    IngestionComplete,
)

from ..request import request_presigned_url, keep_upload_record, storage_upload

logger = logging.getLogger(__name__)


def extract_csv_excel_headers(filename: str, content: bytes) -> list[str]:
    """Extracts column headers from either CSV or XLSX bytes."""
    try: 
        if filename.endswith(".xlsx"):
            excel_file = BytesIO(content)
            wb = load_workbook(excel_file, read_only=True, data_only=True)
            sheet = wb.active
            first_row = next(sheet.iter_rows(values_only=True), [])
            return [str(cell) for cell in first_row if cell is not None]

        elif filename.endswith(".csv"):
            text_stream = StringIO(content.decode("utf-8-sig", errors="ignore"))
            reader = csv.reader(text_stream)
            return next(reader, [])

    except Exception as e:
        logger.error(f"Failed to extract headers for {filename}: {e}")

    return []


async def _load_to_storage(files: list[IngestionCreate]) -> list[IngestionResponse]:
    async with Semaphore(5):
        # Request connections
        async with AsyncClient() as client:
            connection_requests = [
                request_presigned_url(client=client, file=file) for file in files
            ]

            connection_responses: list[IngestionResponse] = await gather(
                *connection_requests
            )

            upload_tasks = [
                storage_upload(client=client, file=file, connection_response=response)
                for file, response in zip(files, connection_responses)
                # if response.status == FileStatus.CREATED
            ]

            results = await gather(*upload_tasks)

    return results


async def _record_file(files: list[IngestionComplete]) -> None:
    async with Semaphore(5):
        async with AsyncClient() as client:
            tasks = [
                create_task(keep_upload_record(client=client, file=file))
                for file in files
            ]

            await gather(*tasks)


__all__ = ["extract_csv_excel_headers", "_load_to_storage", "_record_file"]
